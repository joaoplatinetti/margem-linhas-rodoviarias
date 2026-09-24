"""
A pasta de trabalho formatada, celula por celula, com openpyxl.

Por que openpyxl direto e nao `DataFrame.to_excel`: o destinatario desta saida e
alguem que vai abrir o arquivo e decidir sobre linha de onibus, nao um script.
Isso exige formato de moeda com duas casas, percentual com uma, R$/assento-km com
quatro (e um numero de centavos: com duas casas ele viraria "0,12" em tudo e
perderia a diferenca entre linhas), cabecalho congelado, filtro, semaforo na
classificacao e largura de coluna que nao corte o nome da cidade. `to_excel`
entrega os numeros crus e o trabalho de formatacao ficaria para quem recebe.

Tres decisoes de apresentacao que valem registro:

- **A aba Linhas e a janela de 12 meses, nao o periodo inteiro.** E a aba de
  decisao, e a decisao usa a janela fechada. O periodo completo esta em Mensal.
- **O total da rede aparece no fim da aba Linhas**, em negrito, calculado por
  agregacao e nao por soma de coluna: somar a coluna de RASK daria um numero sem
  significado, e a planilha nao deve nem sugerir isso.
- **Metodologia e Dicionario sao abas, nao um README a parte.** Planilha circula
  desacompanhada; se a premissa nao esta dentro do arquivo, ela se perde no
  primeiro encaminhamento.
"""
from __future__ import annotations

import logging

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from margem import config, custos, indicadores

log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Paleta e estilos
# ---------------------------------------------------------------------------

AZUL_ESCURO = "1F3864"
AZUL_CLARO = "D9E1F2"
CINZA = "F2F2F2"
VERDE = "C6EFCE"
AMARELO = "FFEB9C"
VERMELHO = "FFC7CE"

FONTE_TITULO = Font(bold=True, size=14, color=AZUL_ESCURO)
FONTE_SUBTITULO = Font(bold=True, size=11, color=AZUL_ESCURO)
FONTE_CABECALHO = Font(bold=True, color="FFFFFF", size=10)
PREENCHIMENTO_CABECALHO = PatternFill("solid", fgColor=AZUL_ESCURO)
BORDA_FINA = Border(*[Side(style="thin", color="BFBFBF")] * 4)

CORES_CLASSIFICACAO = {"MANTER": VERDE, "AJUSTAR": AMARELO, "REVER": VERMELHO}

# Formato de numero por coluna. Os codigos seguem a convencao en-US do OOXML
# (`#,##0.00`); o Excel os renderiza com o separador do idioma do usuario, entao
# em pt-BR isso aparece como 1.234,56 — nao e preciso (nem correto) inverter aqui.
MOEDA = 'R$ #,##0.00'
MOEDA_MILHAR = 'R$ #,##0'
UNITARIO = 'R$ #,##0.0000'
PERCENTUAL = '0.0%'
INTEIRO = '#,##0'
DECIMAL = '0.00'

FORMATOS: dict[str, str] = {
    "receita": MOEDA_MILHAR,
    "custo_variavel": MOEDA_MILHAR,
    "custo_total": MOEDA_MILHAR,
    "custo_fixo_rateado": MOEDA_MILHAR,
    "custo_combustivel": MOEDA_MILHAR,
    "custo_pneus": MOEDA_MILHAR,
    "custo_manutencao": MOEDA_MILHAR,
    "custo_tripulacao": MOEDA_MILHAR,
    "custo_pedagio": MOEDA_MILHAR,
    "custo_periodo": MOEDA_MILHAR,
    "margem_contribuicao": MOEDA_MILHAR,
    "resultado": MOEDA_MILHAR,
    "tarifa_media": MOEDA,
    "tripulacao_partida": MOEDA,
    "margem_contribuicao_pct": PERCENTUAL,
    "resultado_pct": PERCENTUAL,
    "load_factor": PERCENTUAL,
    "lf_equilibrio": PERCENTUAL,
    "lf_equilibrio_variavel": PERCENTUAL,
    "folga_lf": PERCENTUAL,
    "ocupacao_base": PERCENTUAL,
    "participacao": PERCENTUAL,
    "yield_pax_km": UNITARIO,
    "rask": UNITARIO,
    "cask_variavel": UNITARIO,
    "cask_total": UNITARIO,
    "spread_rask_cask": UNITARIO,
    "mc_por_ask": UNITARIO,
    "receita_km": MOEDA,
    "variavel_km": MOEDA,
    "combustivel_km": MOEDA,
    "pneus_km": MOEDA,
    "manutencao_km": MOEDA,
    "pedagio_km": MOEDA,
    "tripulacao_km": MOEDA,
    "tarifa_km_referencia": UNITARIO,
    "custo_fixo_km": MOEDA,
    "passageiros": INTEIRO,
    "passageiros_km": INTEIRO,
    "assentos_km": INTEIRO,
    "km_rodados": INTEIRO,
    "partidas": INTEIRO,
    "km": INTEIRO,
    "assentos": INTEIRO,
    "duracao_horas": DECIMAL,
    "fator_tarifa": DECIMAL,
    "amplitude_sazonal": DECIMAL,
    "indice_sazonal": DECIMAL,
}

# Rotulo legivel de cada coluna. O cabecalho da planilha nao deve ser
# `margem_contribuicao_pct`: quem abre o arquivo nao tem por que ler snake_case.
ROTULOS: dict[str, str] = {
    "linha_id": "Cod.",
    "linha": "Linha",
    "origem": "Origem",
    "destino": "Destino",
    "corredor": "Corredor",
    "classe": "Classe",
    "perfil": "Perfil",
    "km": "Km",
    "assentos": "Assentos",
    "frequencia_semanal": "Part./sem.",
    "duracao_horas": "Horas",
    "tripulantes": "Tripul.",
    "ano_mes": "Mes",
    "ano": "Ano",
    "mes": "Mes num.",
    "dias_no_mes": "Dias",
    "indice_sazonal": "Ind. sazonal",
    "partidas": "Partidas",
    "km_rodados": "Km rodados",
    "assentos_km": "ASK (assento-km)",
    "passageiros": "Passageiros",
    "passageiros_km": "RPK (pax-km)",
    "receita": "Receita",
    "tarifa_media": "Tarifa media",
    "custo_combustivel": "Combustivel",
    "custo_pneus": "Pneus",
    "custo_manutencao": "Manutencao",
    "custo_tripulacao": "Tripulacao",
    "custo_pedagio": "Pedagio",
    "custo_variavel": "Custo variavel",
    "custo_fixo_rateado": "Fixo rateado",
    "custo_fixo_km": "Fixo R$/km",
    "custo_total": "Custo total",
    "margem_contribuicao": "Margem contrib.",
    "margem_contribuicao_pct": "MC %",
    "mc_por_ask": "MC / ASK",
    "resultado": "Resultado",
    "resultado_pct": "Resultado %",
    "load_factor": "Load factor",
    "lf_equilibrio": "LF de equilibrio",
    "lf_equilibrio_variavel": "LF equil. (variavel)",
    "folga_lf": "Folga de ocupacao",
    "yield_pax_km": "Yield (R$/pax-km)",
    "rask": "RASK",
    "cask_variavel": "CASK variavel",
    "cask_total": "CASK total",
    "spread_rask_cask": "Spread RASK-CASK",
    "receita_km": "Receita/km",
    "classificacao": "Classificacao",
    "no_limite": "No limite",
    "motivo": "Motivo",
    "alavanca": "Alavanca",
    "componente": "Componente",
    "rotulo": "Componente",
    "grupo": "Grupo",
    "base": "Base",
    "entra_na_mc": "Entra na MC",
    "premissa": "Premissa",
    "custo_periodo": "Custo no periodo",
    "participacao": "% do custo",
    "variavel_km": "Variavel R$/km",
    "combustivel_km": "Combustivel R$/km",
    "pneus_km": "Pneus R$/km",
    "manutencao_km": "Manutencao R$/km",
    "pedagio_km": "Pedagio R$/km",
    "tripulacao_partida": "Tripulacao R$/partida",
    "tripulacao_km": "Tripulacao R$/km",
    "tarifa_km_referencia": "Tarifa ref. R$/pax-km",
    "faixa_distancia": "Faixa de distancia",
    "dupla_tripulacao": "Dupla tripulacao",
}

LARGURA_MINIMA = 9
LARGURA_MAXIMA = 42


# ---------------------------------------------------------------------------
# Blocos de escrita
# ---------------------------------------------------------------------------

def _titulo(ws: Worksheet, texto: str, subtitulo: str = "", linha: int = 1) -> int:
    """Escreve titulo e subtitulo e devolve a proxima linha livre."""
    ws.cell(row=linha, column=1, value=texto).font = FONTE_TITULO
    if subtitulo:
        celula = ws.cell(row=linha + 1, column=1, value=subtitulo)
        celula.font = Font(italic=True, size=9, color="595959")
        return linha + 3
    return linha + 2


def _tabela(
    ws: Worksheet,
    df: pd.DataFrame,
    linha_inicial: int = 1,
    congelar: bool = True,
    filtro: bool = True,
    linha_total: pd.Series | None = None,
) -> int:
    """Escreve o DataFrame como tabela formatada. Devolve a proxima linha livre.

    Faz o trabalho que `to_excel` nao faz: rotulo legivel no cabecalho, formato
    de numero por coluna, semaforo na classificacao, zebrado, borda, filtro e
    painel congelado na primeira coluna de dados.
    """
    colunas = list(df.columns)

    for indice, nome in enumerate(colunas, start=1):
        celula = ws.cell(row=linha_inicial, column=indice, value=ROTULOS.get(nome, nome))
        celula.font = FONTE_CABECALHO
        celula.fill = PREENCHIMENTO_CABECALHO
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = BORDA_FINA

    for deslocamento, (_, registro) in enumerate(df.iterrows(), start=1):
        linha = linha_inicial + deslocamento
        zebrado = PatternFill("solid", fgColor=CINZA) if deslocamento % 2 == 0 else None

        for indice, nome in enumerate(colunas, start=1):
            valor = registro[nome]
            if isinstance(valor, pd.Timestamp):
                valor = valor.date()
            elif pd.api.types.is_bool(valor):
                valor = "sim" if valor else "nao"
            elif pd.isna(valor):
                valor = None

            celula = ws.cell(row=linha, column=indice, value=valor)
            celula.border = BORDA_FINA
            if nome in FORMATOS:
                celula.number_format = FORMATOS[nome]
            if zebrado is not None:
                celula.fill = zebrado

            # Semaforo: pinta a propria celula de classificacao. Formatacao
            # condicional por texto tambem resolveria, mas obrigaria a repetir a
            # lista de rotulos dentro da planilha.
            if nome == "classificacao" and valor in CORES_CLASSIFICACAO:
                celula.fill = PatternFill("solid", fgColor=CORES_CLASSIFICACAO[valor])
                celula.font = Font(bold=True)

    ultima_linha = linha_inicial + len(df)

    if linha_total is not None:
        ultima_linha += 1
        for indice, nome in enumerate(colunas, start=1):
            valor = linha_total.get(nome) if nome in linha_total.index else None
            if indice == 1:
                valor = "REDE"
            celula = ws.cell(row=ultima_linha, column=indice, value=valor)
            celula.font = Font(bold=True)
            celula.fill = PatternFill("solid", fgColor=AZUL_CLARO)
            celula.border = BORDA_FINA
            if nome in FORMATOS:
                celula.number_format = FORMATOS[nome]

    if filtro:
        ws.auto_filter.ref = (
            f"A{linha_inicial}:{get_column_letter(len(colunas))}{linha_inicial + len(df)}"
        )
    if congelar:
        ws.freeze_panes = ws.cell(row=linha_inicial + 1, column=2)

    _ajustar_largura(ws, df, colunas)
    return ultima_linha + 2


def _ajustar_largura(ws: Worksheet, df: pd.DataFrame, colunas: list[str]) -> None:
    """Largura pelo conteudo, limitada nas duas pontas.

    Sem o teto, a coluna de premissa (uma frase inteira) empurra o resto da
    planilha para fora da tela; sem o piso, uma coluna de codigo curto fica
    estreita demais para o proprio cabecalho.
    """
    for indice, nome in enumerate(colunas, start=1):
        rotulo = ROTULOS.get(nome, nome)
        conteudo = df[nome].astype(str).str.len().max() if len(df) else 0
        largura = max(LARGURA_MINIMA, min(LARGURA_MAXIMA, max(len(rotulo) + 3, int(conteudo) + 2)))
        ws.column_dimensions[get_column_letter(indice)].width = largura


def _escala_de_cor(ws: Worksheet, coluna: str, colunas: list[str],
                   primeira: int, ultima: int, invertida: bool = False) -> None:
    """Escala de tres cores na coluna indicada (vermelho -> amarelo -> verde)."""
    if coluna not in colunas:
        return
    letra = get_column_letter(colunas.index(coluna) + 1)
    cores = ["F8696B", "FFEB84", "63BE7B"]
    if invertida:
        cores.reverse()
    ws.conditional_formatting.add(
        f"{letra}{primeira}:{letra}{ultima}",
        ColorScaleRule(
            start_type="min", start_color=cores[0],
            mid_type="percentile", mid_value=50, mid_color=cores[1],
            end_type="max", end_color=cores[2],
        ),
    )


def _barra_de_dados(ws: Worksheet, coluna: str, colunas: list[str],
                    primeira: int, ultima: int) -> None:
    if coluna not in colunas:
        return
    letra = get_column_letter(colunas.index(coluna) + 1)
    ws.conditional_formatting.add(
        f"{letra}{primeira}:{letra}{ultima}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                    color="638EC6", showValue=True),
    )


def _paragrafos(ws: Worksheet, blocos: list[tuple[str, list[str]]], linha: int = 3) -> None:
    """Texto corrido em abas de documentacao, com quebra automatica."""
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 112
    for titulo, paragrafos in blocos:
        celula = ws.cell(row=linha, column=2, value=titulo)
        celula.font = FONTE_SUBTITULO
        linha += 1
        for texto in paragrafos:
            celula = ws.cell(row=linha, column=2, value=texto)
            celula.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[linha].height = 14 * (1 + len(texto) // 95)
            linha += 1
        linha += 1


def _impressao(ws: Worksheet, paisagem: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


# ---------------------------------------------------------------------------
# As abas
# ---------------------------------------------------------------------------

def _aba_resumo(wb: Workbook, fato: pd.DataFrame, janela: pd.DataFrame) -> None:
    ws = wb.active
    ws.title = "Resumo"
    rede = indicadores.resumo_rede(fato)
    meses = janela.attrs.get("meses") or sorted(fato["ano_mes"].unique())

    linha = _titulo(
        ws, "Margem por linha rodoviaria",
        f"Dataset sintetico | {fato['linha_id'].nunique()} linhas | "
        f"{fato['ano_mes'].nunique()} meses ({min(fato['ano_mes'])} a {max(fato['ano_mes'])})",
    )

    ws.cell(row=linha, column=1, value="A rede no periodo completo").font = FONTE_SUBTITULO
    linha += 1
    indicadores_rede = [
        ("Receita", rede["receita"], MOEDA_MILHAR),
        ("Custo variavel", rede["custo_variavel"], MOEDA_MILHAR),
        ("Margem de contribuicao", rede["margem_contribuicao"], MOEDA_MILHAR),
        ("MC %", rede["margem_contribuicao_pct"], PERCENTUAL),
        ("Custo total (com fixo)", rede["custo_total"], MOEDA_MILHAR),
        ("Resultado", rede["resultado"], MOEDA_MILHAR),
        ("Resultado %", rede["resultado_pct"], PERCENTUAL),
        ("Passageiros", rede["passageiros"], INTEIRO),
        ("Load factor", rede["load_factor"], PERCENTUAL),
        ("Yield (R$/pax-km)", rede["yield_pax_km"], UNITARIO),
        ("RASK", rede["rask"], UNITARIO),
        ("CASK variavel", rede["cask_variavel"], UNITARIO),
        ("CASK total", rede["cask_total"], UNITARIO),
        ("Spread RASK-CASK", rede["spread_rask_cask"], UNITARIO),
    ]
    for rotulo, valor, formato in indicadores_rede:
        ws.cell(row=linha, column=1, value=rotulo).font = Font(bold=True, size=10)
        celula = ws.cell(row=linha, column=2, value=float(valor))
        celula.number_format = formato
        linha += 1

    linha += 1
    ws.cell(row=linha, column=1, value="Decisao por linha").font = FONTE_SUBTITULO
    ws.cell(row=linha, column=4,
            value=f"janela de {len(meses)} meses: {meses[0]} a {meses[-1]}").font = Font(
        italic=True, size=9, color="595959")
    linha += 1

    contagem = janela["classificacao"].value_counts().reindex(config.CLASSIFICACOES, fill_value=0)
    leitura = {
        "MANTER": "cobre o custo cheio (RASK >= CASK total)",
        "AJUSTAR": "cobre o variavel, nao o fixo rateado",
        "REVER": "nao cobre nem o custo variavel",
    }
    for nome in config.CLASSIFICACOES:
        celula = ws.cell(row=linha, column=1, value=nome)
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor=CORES_CLASSIFICACAO[nome])
        ws.cell(row=linha, column=2, value=int(contagem[nome])).number_format = INTEIRO
        receita_classe = janela.loc[janela["classificacao"] == nome, "receita"].sum()
        celula = ws.cell(row=linha, column=3, value=float(receita_classe))
        celula.number_format = MOEDA_MILHAR
        ws.cell(row=linha, column=4, value=leitura[nome]).font = Font(size=9, color="595959")
        linha += 1

    linha += 1
    ws.cell(row=linha, column=1, value="As linhas em risco").font = FONTE_SUBTITULO
    linha += 1
    risco = janela[janela["classificacao"] != "MANTER"].sort_values("spread_rask_cask")
    colunas_risco = ["linha_id", "linha", "km", "classe", "load_factor",
                     "lf_equilibrio", "folga_lf", "yield_pax_km", "rask",
                     "cask_total", "spread_rask_cask", "margem_contribuicao_pct",
                     "classificacao", "alavanca"]
    _tabela(ws, risco[colunas_risco], linha_inicial=linha, congelar=False, filtro=False)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 40
    _impressao(ws, paisagem=False)


def _aba_linhas(wb: Workbook, janela: pd.DataFrame) -> None:
    ws = wb.create_sheet("Linhas")
    meses = janela.attrs.get("meses") or []
    periodo = f"{meses[0]} a {meses[-1]}" if meses else ""

    colunas = [
        "linha_id", "linha", "corredor", "classe", "perfil", "km", "assentos",
        "frequencia_semanal", "partidas", "assentos_km", "passageiros",
        "passageiros_km", "load_factor", "lf_equilibrio", "folga_lf",
        "tarifa_media", "yield_pax_km", "rask",
        "cask_variavel", "cask_total", "spread_rask_cask", "receita",
        "custo_variavel", "margem_contribuicao", "margem_contribuicao_pct",
        "custo_fixo_rateado", "custo_total", "resultado", "classificacao",
        "no_limite", "motivo", "alavanca",
    ]
    tabela = janela[colunas]

    linha = _titulo(
        ws, "Decisao por linha",
        f"Janela de decisao: {periodo}. Razoes calculadas sobre as somas da janela, "
        "nunca como media dos meses.",
    )
    primeira_dados = linha + 1
    # A janela ja vem com as aditivas somadas, km_rodados incluido.
    total = indicadores.agregar(janela).iloc[0]
    _tabela(ws, tabela, linha_inicial=linha, linha_total=total)

    ultima = linha + len(tabela)
    _escala_de_cor(ws, "margem_contribuicao_pct", colunas, primeira_dados, ultima)
    _escala_de_cor(ws, "spread_rask_cask", colunas, primeira_dados, ultima)
    _escala_de_cor(ws, "folga_lf", colunas, primeira_dados, ultima)
    _barra_de_dados(ws, "load_factor", colunas, primeira_dados, ultima)
    _impressao(ws)


def _aba_mensal(wb: Workbook, fato: pd.DataFrame) -> None:
    ws = wb.create_sheet("Mensal")
    colunas = [
        "linha_id", "linha", "ano_mes", "dias_no_mes", "partidas", "km_rodados",
        "assentos_km", "passageiros", "passageiros_km", "load_factor",
        "tarifa_media", "yield_pax_km", "rask", "receita",
        "custo_combustivel", "custo_pneus", "custo_manutencao",
        "custo_tripulacao", "custo_pedagio", "custo_variavel",
        "custo_fixo_rateado", "custo_total", "margem_contribuicao",
        "margem_contribuicao_pct", "cask_variavel", "cask_total",
        "spread_rask_cask", "resultado",
    ]
    tabela = fato[colunas].sort_values(["linha_id", "ano_mes"], ignore_index=True)
    linha = _titulo(
        ws, "Fato mensal",
        "Um registro por linha e mes. E a base de tudo: as outras abas sao "
        "agregacoes desta.",
    )
    _tabela(ws, tabela, linha_inicial=linha)
    _impressao(ws)


def _aba_custos(wb: Workbook, fato: pd.DataFrame, catalogo: pd.DataFrame) -> None:
    ws = wb.create_sheet("Custos")

    linha = _titulo(
        ws, "Modelo de custo",
        "Cada componente com a sua base e a sua premissa. So os de base variavel "
        "entram na margem de contribuicao.",
    )

    premissas = custos.premissas(fato)[
        ["rotulo", "grupo", "base", "entra_na_mc", "premissa", "custo_periodo", "participacao"]
    ]
    linha = _tabela(ws, premissas, linha_inicial=linha, congelar=False, filtro=False)

    ws.cell(row=linha, column=1, value="Blocos de custo fixo mensal").font = FONTE_SUBTITULO
    linha += 1
    for nome, valor in config.CUSTO_FIXO_MENSAL.items():
        ws.cell(row=linha, column=1, value=nome.replace("_", " ").capitalize())
        ws.cell(row=linha, column=2, value=valor).number_format = MOEDA_MILHAR
        linha += 1
    ws.cell(row=linha, column=1, value="Total mensal").font = Font(bold=True)
    celula = ws.cell(row=linha, column=2, value=sum(config.CUSTO_FIXO_MENSAL.values()))
    celula.number_format = MOEDA_MILHAR
    celula.font = Font(bold=True)
    linha += 1
    fixo_km = fato.groupby("ano_mes")["custo_fixo_km"].first()
    ws.cell(row=linha, column=1, value="Rateio medio (R$/km rodado)")
    ws.cell(row=linha, column=2, value=float(fixo_km.mean())).number_format = MOEDA
    linha += 3

    ws.cell(row=linha, column=1, value="Custo variavel unitario por linha").font = FONTE_SUBTITULO
    ws.cell(row=linha, column=6,
            value="Repare no degrau da tripulacao acima de 600 km: entra o segundo "
                  "motorista e o custo por km quase dobra.").font = Font(
        italic=True, size=9, color="595959")
    linha += 1
    unitario = custos.custo_unitario_por_linha(catalogo).sort_values("km")
    _tabela(ws, unitario, linha_inicial=linha, congelar=False, filtro=False)
    _impressao(ws)


def _aba_dicionario(wb: Workbook) -> None:
    ws = wb.create_sheet("Dicionario")
    dicionario = pd.DataFrame(
        [
            ("ASK (assento-km)", "assentos x km x partidas", "assento-km",
             "A oferta. Denominador de RASK, CASK e load factor."),
            ("RPK (pax-km)", "passageiros x km", "passageiro-km",
             "A demanda transportada. Aqui todo passageiro viaja o trecho inteiro."),
            ("Load factor", "RPK / ASK", "%",
             "Quanto da oferta foi vendida. Teto do modelo: 96%."),
            ("Tarifa media", "receita / passageiros", "R$",
             "O bilhete medio. Sobe no pico, quando ha menos promocional."),
            ("Yield", "receita / RPK", "R$ por pax-km",
             "Preco por unidade de distancia vendida. Comparavel entre linhas de "
             "tamanhos diferentes, ao contrario da tarifa media."),
            ("RASK", "receita / ASK", "R$ por assento-km",
             "Receita por assento oferecido. Identidade: RASK = yield x load factor."),
            ("CASK variavel", "custo variavel / ASK", "R$ por assento-km",
             "Custo que so existe se o onibus sair."),
            ("CASK total", "(variavel + fixo rateado) / ASK", "R$ por assento-km",
             "Custo cheio, com o pedaco de estrutura rateado por km rodado."),
            ("Margem de contribuicao", "receita - custo variavel", "R$",
             "Quanto a linha sobra para pagar o fixo. Negativa: cada partida "
             "destroi caixa."),
            ("MC %", "MC / receita", "%",
             "A margem em percentual da receita, comparavel entre linhas."),
            ("LF de equilibrio", "CASK total / yield", "%",
             "A ocupacao que a linha precisaria ter, ao yield que ja pratica, para "
             "empatar com o custo cheio. Sai da identidade RASK = yield x LF."),
            ("Folga de ocupacao", "load factor - LF de equilibrio", "p.p.",
             "O mesmo que o spread, em unidade que a operacao entende: 'faltam 8 "
             "pontos de ocupacao' e acionavel, '-R$ 0,015 por assento-km' nao."),
            ("Spread RASK-CASK", "RASK - CASK total", "R$ por assento-km",
             "O resultado por assento oferecido. Positivo: a linha cobre o custo "
             "cheio. E o que classifica."),
            ("Resultado", "receita - custo total", "R$",
             "O mesmo spread, em reais."),
            ("MANTER", "spread >= 0", "-",
             "Cobre o custo cheio. Acompanhar."),
            ("AJUSTAR", "MC > 0 e spread < 0", "-",
             "Contribui para o fixo mas nao o cobre. Mexer em frequencia, classe, "
             "horario ou preco — nao na existencia da linha."),
            ("REVER", "MC <= 0", "-",
             "Nao cobre nem o variavel. Volume nao resolve: vender mais assento "
             "com MC negativa piora o resultado."),
            ("No limite", f"|spread| < R$ {config.SPREAD_NO_LIMITE:.3f}/ASK", "-",
             "Qualifica a classificacao: cobre o custo, mas uma alta de diesel "
             "reverte."),
        ],
        columns=["indicador", "formula", "unidade", "como_ler"],
    )
    dicionario.columns = ["Indicador", "Formula", "Unidade", "Como ler"]

    linha = _titulo(ws, "Dicionario de indicadores",
                    "As tres identidades: RASK = yield x load factor | "
                    "spread = RASK - CASK total | MC = receita - custo variavel")
    _tabela(ws, dicionario, linha_inicial=linha, congelar=False, filtro=False)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 72
    for celulas in ws.iter_rows(min_row=linha + 1, min_col=4, max_col=4):
        for celula in celulas:
            celula.alignment = Alignment(wrap_text=True, vertical="top")
    _impressao(ws, paisagem=False)


def _aba_metodologia(wb: Workbook, fato: pd.DataFrame) -> None:
    ws = wb.create_sheet("Metodologia")
    fixo_total = sum(config.CUSTO_FIXO_MENSAL.values())

    blocos: list[tuple[str, list[str]]] = [
        ("O que este arquivo e", [
            "Um exercicio de analise de margem por linha de onibus interestadual. "
            "As cidades e as distancias sao reais; todos os numeros de operacao, "
            "demanda, tarifa e custo sao SINTETICOS, gerados por codigo a partir "
            "de premissas declaradas. Nada aqui vem de sistema, planilha ou base "
            "de empresa nenhuma.",
            f"Periodo: {min(fato['ano_mes'])} a {max(fato['ano_mes'])} "
            f"({fato['ano_mes'].nunique()} meses). "
            f"Escopo: {fato['linha_id'].nunique()} linhas. "
            f"Semente do gerador: {config.SEED} — a mesma execucao produz sempre "
            "os mesmos numeros.",
        ]),
        ("Como a demanda foi gerada", [
            "passageiros = ocupacao base x sazonalidade^amplitude x tendencia x "
            "ruido, cortado no teto de oferta.",
            "A sazonalidade e um indice mensal de media 1 (julho 1,22 e janeiro "
            "1,18 no topo; marco 0,87 e maio 0,88 no fundo), elevado a um "
            "expoente proprio de cada linha: linha corporativa achata a curva "
            "(Goiania-Brasilia, 0,4), linha de lazer exagera (Sao Paulo-Foz, 1,6). "
            "Sazonalidade identica em todas as linhas e a assinatura mais obvia "
            "de dado inventado.",
            f"O teto de ocupacao e {config.OCUPACAO_MAXIMA:.0%}: nenhuma linha "
            "vende o mes inteiro, sobra assento no horario ruim e no dia de "
            "semana. Sem esse teto, o pico de julho produziria load factor acima "
            "de 100%.",
            f"A tendencia e de {config.TENDENCIA_ANUAL:.0%} ao ano, centrada no "
            "meio do periodo, e o ruido e lognormal com desvio de "
            f"{config.RUIDO_DEMANDA:.0%} ao mes.",
        ]),
        ("Como a tarifa foi montada", [
            "A tarifa parte de um valor de referencia em R$ por passageiro-km da "
            "classe (convencional 0,175 ate leito 0,400) e cai com a distancia "
            "pelo fator (500/km)^0,18 — o custo de terminal, embarque e venda se "
            "dilui em percurso maior, e por isso R$/km de uma linha de 200 km e "
            "naturalmente maior que o de um leito de 1.000 km.",
            "Depois entra o fator de pressao competitiva da linha (0,92 em "
            "Sao Paulo-Rio, onde a aerea concorre; 1,02 em Sao Paulo-Foz, rota "
            "turistica com pouca alternativa direta) e a sensibilidade da tarifa "
            "ao mes: no pico sobra menos promocional e a tarifa media sobe.",
        ]),
        ("O modelo de custo", [
            "Cinco componentes variaveis, com a base explicita. Por km: "
            "combustivel (diesel dividido pelo consumo da classe), pneus, "
            "manutencao (com fator de agressividade do corredor) e pedagio (R$/km "
            "do corredor). Por partida: tripulacao, que e jornada x custo-hora "
            "mais diaria e pernoite.",
            "Tripulacao por PARTIDA e nao por km e a decisao que mais muda o "
            "resultado. Acima de 600 km entra o segundo motorista e o custo de "
            "tripulacao por km quase dobra — uma linha de 610 km carrega quase o "
            "dobro de tripulacao por km de uma de 590 km, com receita por km "
            "parecida. E o degrau que explica boa parte das linhas classificadas "
            "como AJUSTAR.",
            f"O fixo e um bloco mensal de R$ {fixo_total:,.0f} ".replace(",", ".")
            + "(garagem, administracao e vendas, depreciacao e seguros), rateado "
            "pelos km rodados do mes. A soma do rateado fecha com o total "
            "declarado em cada mes.",
        ]),
        ("A decisao", [
            "A classificacao usa a janela dos 12 ultimos meses somados, nunca o "
            "mes isolado: uma linha de lazer cobre o custo cheio com folga em "
            "julho e nao cobre nem o variavel em maio, e decidir mes a mes faria "
            "a mesma linha alternar entre MANTER e REVER tres vezes ao ano.",
            "A escada tem dois degraus. MANTER cobre o custo cheio. AJUSTAR cobre "
            "o variavel mas nao o fixo rateado — e a linha CONTRIBUI: tirada da "
            "malha, o fixo que ela cobria se redistribui entre as outras e piora "
            "todas. REVER nao cobre nem o variavel, e ai volume nao resolve.",
            "A coluna Alavanca vem da identidade RASK = yield x load factor: se a "
            "linha nao cobre o custo, falta preco ou falta ocupacao, e as duas "
            "coisas pedem acao diferente. A comparacao e contra a mediana da rede.",
        ]),
        ("Limites declarados", [
            "Todo passageiro viaja o trecho inteiro, entao passageiro-km e "
            "simplesmente passageiros x km. Um modelo com secoes intermediarias "
            "exigiria a ocupacao trecho a trecho para ratear a capacidade, e sem "
            "isso qualquer 'ocupacao por secao' seria um rateio inventado.",
            "A linha e medida em um sentido. Sentido de volta com ocupacao "
            "diferente (o classico vazio do retorno de feriado) nao esta modelado.",
            "O rateio do fixo por km e uma convencao, nao uma verdade: ele segue a "
            "intensidade de uso de frota e oficina, mas nenhuma decisao de cortar "
            "linha deve sair so dele. E por isso que a margem de contribuicao "
            "aparece sempre ao lado.",
        ]),
    ]

    _titulo(ws, "Metodologia", "Como o dataset foi gerado e como os indicadores "
                               "foram calculados.")
    _paragrafos(ws, blocos, linha=4)
    _impressao(ws, paisagem=False)


# ---------------------------------------------------------------------------
# Montagem
# ---------------------------------------------------------------------------

def _propriedades(wb: Workbook, fato: pd.DataFrame) -> None:
    """Metadados fixos, para o arquivo versionado nao mudar a cada execucao.

    O openpyxl carimba a data de criacao no momento do save. Como a saida e
    versionada no repo, isso faria `margem tudo` produzir um arquivo diferente
    todo dia sem uma unica celula ter mudado — e o `git status` sujo deixaria de
    significar qualquer coisa. A data usada e a do ultimo mes do dataset, que e
    a informacao que o arquivo de fato carrega.

    Os carimbos internos do zip continuam variando (e limitacao do openpyxl),
    entao os bytes ainda mudam; o que isto resolve e a metadata visivel em
    Arquivo > Informacoes.
    """
    ultimo = pd.Period(max(fato["ano_mes"]), freq="M").to_timestamp(how="end")
    wb.properties.creator = "margem-linhas-rodoviarias"
    wb.properties.lastModifiedBy = "margem-linhas-rodoviarias"
    wb.properties.title = "Margem por linha rodoviaria"
    wb.properties.description = (
        "Dataset sintetico. Cidades e distancias reais; operacao, demanda, "
        "tarifa e custo gerados por codigo a partir de premissas declaradas."
    )
    wb.properties.created = ultimo
    wb.properties.modified = ultimo


def exportar(fato: pd.DataFrame, janela: pd.DataFrame, catalogo: pd.DataFrame) -> str:
    """Monta e grava a pasta de trabalho. Devolve o caminho."""
    wb = Workbook()
    _propriedades(wb, fato)
    _aba_resumo(wb, fato, janela)
    _aba_linhas(wb, janela)
    _aba_mensal(wb, fato)
    _aba_custos(wb, fato, catalogo)
    _aba_dicionario(wb)
    _aba_metodologia(wb, fato)

    config.SAIDA.mkdir(parents=True, exist_ok=True)
    wb.save(config.EXCEL)
    log.info("planilha: %s (%d abas)", config.EXCEL.name, len(wb.sheetnames))
    return str(config.EXCEL)
